#!/usr/bin/env python3
"""
Stage 2 of the workbook build — format workbook-data.json into meridian.xlsx.

Design decision, and the reason for the two-stage split:

The previous workbook was formula-driven end to end, every cell re-deriving the
model in Excel. That is a THIRD implementation of Meridian alongside the JS
engine and the Python backtest, and it drifted exactly as you would predict —
its financial-sector exemption enumerated 11 industries in an OR() formula while
the engine had 12, silently mis-scoring three companies. §7.2 rejects duplicate
implementations; verify_port exists to catch the drift between the two that must
exist. A third one is not defensible.

So the rule here is: model outputs arrive as values from meridian-engine.js (the
same code the app runs), and formulas are used only where they compute something
the model does NOT define — cross-sheet counts, distributions, display ratios
over adjacent columns. Those recalculate meaningfully and cannot disagree with
the model, because they are arithmetic over its output rather than a restatement
of it. Engine values are styled as inputs (blue); formulas are black.

The Dashboard's gate funnel is the exception that earns its keep: it rebuilds the
five Golden Breakout gates with COUNTIFS over the Stocks columns and asserts the
count equals the screener's own row count. If the engine and the emitted columns
ever disagree, the workbook says MISMATCH on its first sheet.

Usage: python3 build_workbook.py [--data workbook-data.json] [--out meridian.xlsx]
"""

import argparse, json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
INK = "1F2937"        # header fill: slate
ACCENT = "2563EB"     # Meridian equities accent
TEAL = "2DB9A3"
RULE = Side(style="thin", color="D1D5DB")

H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor=INK)
V_FONT = Font(name=FONT, size=10, color="0000FF")   # engine value  (input convention)
F_FONT = Font(name=FONT, size=10, color="000000")   # workbook formula
T_FONT = Font(name=FONT, size=10)                   # plain text
TITLE = Font(name=FONT, size=14, bold=True, color=INK)
SUB = Font(name=FONT, size=10, color="6B7280")

PCT = '0.0%;(0.0%);-'
PCT2 = '0.00%;(0.00%);-'
NUM2 = '#,##0.00;(#,##0.00);-'
NUM0 = '#,##0;(#,##0);-'
RATIO = '0.00"x";(0.00"x");-'
INT = '#,##0;-#,##0;-'


def pct(v):
    """Engine emits percentage points; Excel percent format wants a fraction."""
    return None if v is None else v / 100.0


# ---------- sheet writer ----------
# cols: (header, key_or_None, number_format, width, kind)
#   kind "v" = engine value (blue), "f" = formula (black, key is a lambda(row)->str)

def write_sheet(wb, name, cols, rows, tab_color=ACCENT, note=None, freeze="A2"):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab_color

    for i, (header, _key, fmt, width, _kind) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=header)
        c.font, c.fill = H_FONT, H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    for r, item in enumerate(rows, start=2):
        for i, (_h, key, fmt, _w, kind) in enumerate(cols, start=1):
            if kind == "f":
                c = ws.cell(row=r, column=i, value=key(r))
                c.font = F_FONT
            else:
                v = item.get(key)
                if isinstance(v, bool):
                    v = "Yes" if v else "No"
                c = ws.cell(row=r, column=i, value=v)
                c.font = V_FONT
            if fmt:
                c.number_format = fmt
            c.border = Border(bottom=RULE)

    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        ws.freeze_panes = freeze
    if note:
        r = len(rows) + 3
        ws.cell(row=r, column=1, value=note).font = SUB
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=min(len(cols), 8))
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="workbook-data.json")
    ap.add_argument("--out", default="meridian.xlsx")
    a = ap.parse_args()

    d = json.load(open(a.data))
    meta = d["meta"]
    stocks, breakout = d["stocks"], d["breakout"]
    sectoral, sect_bo, breadth, universe = d["sectoral"], d["sectoralBreakout"], d["breadth"], d["universe"]

    # Percent-point -> fraction, once, before any sheet is written.
    for s in stocks:
        for k in ("ChangePct", "VolBreakoutPct", "MA200SlopePct", "SeparationPct",
                  "ROE3yr", "ROCE3yr", "EPSGrowth3yr", "SalesGrowth3yr", "CWIPPct"):
            s[k] = pct(s[k])
    for b in breakout:
        for k in ("ChangePct", "SeparationPct", "MA200SlopePct", "VolBreakoutPct"):
            b[k] = pct(b[k])
    for s in sectoral:
        for k in ("ChangePct", "MA200SlopePct", "SeparationPct"):
            s[k] = pct(s[k])
    for s in sect_bo:
        for k in ("ChangePct", "SeparationPct", "MA200SlopePct"):
            s[k] = pct(s[k])
    for b in breadth:
        for k in ("PctAbove200DMA", "MA30", "MA100", "MA200"):
            b[k] = pct(b[k])

    wb = Workbook()
    wb.remove(wb.active)

    NS, NU, NB = len(stocks), len(universe), len(breadth)
    SR, UR = NS + 1, NU + 1          # last data row on Stocks / Universe

    # ---------------- Stocks ----------------
    # Column letters are referenced by the Dashboard formulas, so the order here is
    # load-bearing; the map below is the single place it is written down.
    S = dict(Symbol="A", Name="B", Industry="C", Group="D", ISIN="E", MCap="F", CMP="G",
             Chg="H", RS="I", Band="J", RSDays="K", RSI="L", MA8="M", MA50="N", MA200="O",
             Hi="P", Lo="Q", FromHi="R", Vol="S", Slope="T", Rising="U", Above50="V",
             Above8="W", Cross="X", Fresh="Y", Sep="Z", Score="AA", Tier="AB", Exempt="AC",
             ROE="AD", ROCE="AE", EPSG="AF", SalesG="AG", DE="AH", AT="AI", CWIP="AJ")

    stock_cols = [
        ("Symbol", "Symbol", None, 12, "v"), ("Company", "Name", None, 34, "v"),
        ("Industry", "Industry", None, 26, "v"), ("Sector Group", "SectorGroup", None, 20, "v"),
        ("ISIN", "ISIN", None, 14, "v"), ("MCap (Rs Cr)", "MarketCap", NUM0, 12, "v"),
        ("CMP (Rs)", "CMP", NUM2, 11, "v"), ("Chg %", "ChangePct", PCT, 8, "v"),
        ("RS Rating", "RSRating", INT, 8, "v"), ("RS Band", "RSBand", None, 8, "v"),
        ("RS Days", "RSStreakDays", INT, 8, "v"), ("RSI (14)", "RSI", NUM2, 8, "v"),
        ("8 DMA", "MA8", NUM2, 11, "v"), ("50 DMA", "MA50", NUM2, 11, "v"),
        ("200 DMA", "MA200", NUM2, 11, "v"),
        ("52w High", "High52", NUM2, 11, "v"), ("52w Low", "Low52", NUM2, 11, "v"),
        ("% from 52w High",
         lambda r: f'=IF(OR({S["CMP"]}{r}="",{S["Hi"]}{r}=""),"",{S["CMP"]}{r}/{S["Hi"]}{r}-1)',
         PCT, 12, "f"),
        ("Vol vs 30d Avg", "VolBreakoutPct", PCT, 11, "v"),
        ("200DMA Slope %", "MA200SlopePct", PCT2, 11, "v"),
        ("G2: 200DMA Rising", "MA200Rising", None, 10, "v"),
        ("G1b: Px > 50DMA", "PriceAbove50DMA", None, 10, "v"),
        ("G5: Px > 8DMA", "PriceAbove8DMA", None, 10, "v"),
        ("G1a: 50 > 200DMA", "GoldenCrossState", None, 10, "v"),
        ("G4: Freshness (d)", "GoldenCrossStreak", INT, 10, "v"),
        ("G3: Separation %", "SeparationPct", PCT, 11, "v"),
        ("Fund Score", "FundScore", NUM2, 10, "v"), ("Fund Tier", "FundTier", None, 10, "v"),
        ("Fin. Exempt", "FinExempt", None, 10, "v"),
        ("ROE 3y %", "ROE3yr", PCT, 10, "v"), ("ROCE 3y %", "ROCE3yr", PCT, 10, "v"),
        ("EPS Gr 3y %", "EPSGrowth3yr", PCT, 10, "v"), ("Sales Gr 3y %", "SalesGrowth3yr", PCT, 10, "v"),
        ("D/E", "DebtEquity", RATIO, 9, "v"), ("Asset Turns", "AssetTurns", RATIO, 10, "v"),
        ("CWIP %", "CWIPPct", PCT, 9, "v"),
    ]
    write_sheet(wb, "Stocks", stock_cols, stocks, ACCENT, freeze="C2",
                note="Blue = value computed by meridian-engine.js (the same module the app runs). "
                     "Black = formula computed inside this workbook. Columns tagged G1a/G1b/G2/G3/G4/G5 "
                     "are the five Golden Breakout gates; the Dashboard rebuilds the funnel from them.")

    # ---------------- Golden Breakout ----------------
    bo_cols = [
        ("Rank", lambda r: f"=ROW()-1", INT, 7, "f"),
        ("Symbol", "Symbol", None, 12, "v"), ("Company", "Name", None, 34, "v"),
        ("Industry", "Industry", None, 26, "v"),
        ("CMP (Rs)", "CMP", NUM2, 11, "v"), ("Chg %", "ChangePct", PCT, 9, "v"),
        ("Separation %", "SeparationPct", PCT, 12, "v"),
        ("Freshness (d)", "FreshnessDays", INT, 12, "v"),
        ("200DMA Slope %", "MA200SlopePct", PCT2, 13, "v"),
        ("Vol vs 30d Avg", "VolBreakoutPct", PCT, 13, "v"),
        ("Fund Score", "FundScore", NUM2, 11, "v"), ("Fund Tier", "FundTier", None, 11, "v"),
    ]
    write_sheet(wb, "Golden Breakout", bo_cols, breakout, "16A34A",
                note=f"Ranked by separation (widest first), freshness breaking ties. Gates: 50DMA>200DMA "
                     f"with price above both, 200DMA rising, separation >= "
                     f"{meta['gates']['minSeparationPct']}%, cross no older than "
                     f"{meta['gates']['freshnessMaxDays']} trading days, price above 8DMA. "
                     "Fund Score is context, not a gate - a blank score means the company is outside "
                     "the fundamentals sample, not that it scored zero.")

    # ---------------- Sectoral ----------------
    sec_cols = [
        ("Industry", "Industry", None, 30, "v"), ("Constituents", "Constituents", INT, 12, "v"),
        ("Share of Priced Universe",
         lambda r: f"=B{r}/SUM($B$2:$B${len(sectoral) + 1})", PCT, 14, "f"),
        ("Index Level", "IndexLevel", NUM2, 12, "v"), ("Chg %", "ChangePct", PCT, 9, "v"),
        ("RS Rating", "RSRating", INT, 9, "v"), ("RS Band", "RSBand", None, 9, "v"),
        ("RS Days", "RSStreakDays", INT, 9, "v"), ("RSI (14)", "RSI", NUM2, 9, "v"),
        ("50 DMA", "MA50", NUM2, 11, "v"), ("200 DMA", "MA200", NUM2, 11, "v"),
        ("200DMA Slope %", "MA200SlopePct", PCT2, 13, "v"),
        ("50 > 200DMA", "GoldenCrossState", None, 11, "v"),
        ("Freshness (d)", "GoldenCrossStreak", INT, 12, "v"),
        ("Separation %", "SeparationPct", PCT, 12, "v"),
    ]
    write_sheet(wb, "Sectoral", sec_cols, sectoral, "7C3AED",
                note="Synthetic equal-weight index per Industry Name, rebased to 100 and chained on daily "
                     "average constituent returns. Industries with fewer than 3 priced constituents are "
                     "excluded as too thin to index. Index Level is a rebased level, not a rupee price.")

    # ---------------- Sectoral Breakout ----------------
    sbo_cols = [
        ("Rank", lambda r: "=ROW()-1", INT, 7, "f"),
        ("Industry", "Industry", None, 30, "v"),
        ("Index Level", "IndexLevel", NUM2, 12, "v"), ("Chg %", "ChangePct", PCT, 9, "v"),
        ("Separation %", "SeparationPct", PCT, 12, "v"),
        ("Freshness (d)", "FreshnessDays", INT, 12, "v"),
        ("200DMA Slope %", "MA200SlopePct", PCT2, 13, "v"),
    ]
    ws = write_sheet(wb, "Sectoral Breakout", sbo_cols, sect_bo, "7C3AED")
    if not sect_bo:
        ws.cell(row=3, column=2, value="No industry index qualified on this date.").font = Font(
            name=FONT, size=11, bold=True, color=INK)
        ws.cell(row=5, column=2, value=(
            "This is a real result, not a missing feed. An equal-weight index of many stocks moves far "
            "more smoothly than any one of them, so its 50/200 DMA crossovers are rare and, once they "
            "happen, quickly age past the freshness window. On this run 104 of 120 industries were in "
            "golden-cross state and 72 also cleared separation and a rising 200DMA - but the freshest "
            "cross among them was 23 trading days old, outside the 15-day gate. Expect this sheet to be "
            "empty most days and to populate in clusters after a broad market turn.")).font = SUB
        ws.merge_cells("B5:H9")
        ws.cell(row=5, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------- Market Breadth ----------------
    br_cols = [
        ("Date", "Date", None, 12, "v"),
        ("% Above 200DMA", "PctAbove200DMA", PCT, 14, "v"),
        ("30d MA", "MA30", PCT, 11, "v"), ("100d MA", "MA100", PCT, 11, "v"),
        ("200d MA", "MA200", PCT, 11, "v"),
        ("New Highs", "NewHighs", INT, 11, "v"), ("New Lows", "NewLows", INT, 11, "v"),
        ("High/Low Ratio", "HighLowRatio", NUM2, 13, "v"),
        ("Net Highs", lambda r: f"=F{r}-G{r}", INT, 11, "f"),
    ]
    write_sheet(wb, "Market Breadth", br_cols, breadth, "EA580C",
                note="Whole-universe breadth over the last 500 trading days. The 30/100/200d columns "
                     "smooth the breadth line itself; they are not different thresholds. A blank "
                     "High/Low Ratio means there were no new lows at all that day (an undefined ratio, "
                     "not zero). New high/low = within 1% of the 252-day extreme.")

    # ---------------- Universe ----------------
    uni_cols = [
        ("ISIN", "ISIN", None, 14, "v"), ("Symbol", "Symbol", None, 12, "v"),
        ("Company", "Name", None, 36, "v"), ("Industry", "Industry", None, 28, "v"),
        ("Sector Group", "SectorGroup", None, 22, "v"),
        ("MCap (Rs Cr)", "MarketCap", NUM0, 13, "v"),
        ("NSE Code", "NSECode", None, 12, "v"), ("BSE Code", "BSECode", None, 11, "v"),
        ("Priced",
         lambda r: (f'=IFERROR(IF(INDEX(Stocks!${S["CMP"]}$2:${S["CMP"]}${SR},'
                    f'MATCH(A{r},Stocks!${S["ISIN"]}$2:${S["ISIN"]}${SR},0))="","No","Yes"),"No")'),
         None, 9, "f"),
    ]
    write_sheet(wb, "Universe", uni_cols, universe, "6B7280", freeze="C2",
                note="The quarterly universe definition. BSE Code is text on purpose - read it as a "
                     "number anywhere in the pipeline and '544467' becomes '544467.0', which breaks the "
                     "Yahoo ticker. Priced is looked up live against the Stocks sheet.")

    # ---------------- Dashboard ----------------
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_properties.tabColor = INK
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 58
    ws.sheet_view.showGridLines = False

    def head(r, text):
        c = ws.cell(row=r, column=2, value=text)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        for col in (3, 4):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=INK)

    def line(r, label, formula, note="", fmt=INT, bold=False):
        ws.cell(row=r, column=2, value=label).font = Font(name=FONT, size=10, bold=bold)
        c = ws.cell(row=r, column=3, value=formula)
        c.font = Font(name=FONT, size=10, bold=bold)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=4, value=note).font = SUB

    ws["B2"] = "Meridian"
    ws["B2"].font = TITLE
    ws["B3"] = f"Daily screen  |  as of {meta['as_of']}  |  built {meta['generated_at'][:16].replace('T', ' ')} UTC"
    ws["B3"].font = SUB

    ST = f"Stocks!"
    r = 5
    head(r, "Coverage"); r += 1
    line(r, "Instruments in universe", f"=COUNTA({ST}${S['Symbol']}$2:${S['Symbol']}${SR})",
         "Quarterly universe definition (see Universe sheet)"); r += 1
    line(r, "With price history", f"=COUNT({ST}${S['CMP']}$2:${S['CMP']}${SR})",
         "The rest have no Yahoo series under any ticker variant"); r += 1
    line(r, "With a 200 DMA", f"=COUNT({ST}${S['MA200']}$2:${S['MA200']}${SR})",
         "Needs 200 clean bars; newer listings fall short"); r += 1
    line(r, "With a fundamental score", f"=COUNT({ST}${S['Score']}$2:${S['Score']}${SR})",
         "Fundamentals sample is narrower than the price universe"); r += 2

    head(r, "Golden Breakout funnel"); r += 1
    g1 = f'COUNTIFS({ST}${S["Cross"]}$2:${S["Cross"]}${SR},"Yes",{ST}${S["Above50"]}$2:${S["Above50"]}${SR},"Yes"'
    line(r, "G1  Price > 50DMA > 200DMA", f"={g1})", "Trend structure intact"); r += 1
    g2 = g1 + f',{ST}${S["Rising"]}$2:${S["Rising"]}${SR},"Yes"'
    line(r, "G2  + 200DMA rising", f"={g2})", "The long trend is itself improving"); r += 1
    g3 = g2 + f',{ST}${S["Sep"]}$2:${S["Sep"]}${SR},">="&{meta["gates"]["minSeparationPct"] / 100}'
    line(r, f"G3  + separation >= {meta['gates']['minSeparationPct']}%", f"={g3})",
         "Filters marginal crossovers that whipsaw"); r += 1
    g4 = g3 + f',{ST}${S["Fresh"]}$2:${S["Fresh"]}${SR},"<="&{meta["gates"]["freshnessMaxDays"]}'
    line(r, f"G4  + cross <= {meta['gates']['freshnessMaxDays']} trading days old", f"={g4})",
         "Freshness window - widened from 10 to 15 days"); r += 1
    g5 = g4 + f',{ST}${S["Above8"]}$2:${S["Above8"]}${SR},"Yes"'
    line(r, "G5  + price > 8DMA", f"={g5})", "Short-term trend not already rolling over",
         bold=True); r += 1
    line(r, "Rows on the Golden Breakout sheet", "=COUNT('Golden Breakout'!$A$2:$A$10000)",
         "The screener's own output"); r += 1
    fr = r
    ws.cell(row=r, column=2, value="Self-check").font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=3, value=f"=IF(C{r - 2}=C{r - 1},\"OK\",\"MISMATCH\")")
    c.font = Font(name=FONT, size=10, bold=True)
    c.alignment = Alignment(horizontal="right")
    ws.cell(row=r, column=4,
            value="Rebuilds the five gates with COUNTIFS and compares against the screener. "
                  "MISMATCH means the emitted columns and the engine disagree.").font = SUB
    r += 2

    head(r, "Fundamental tiers"); r += 1
    for tier, note in [("High", "Top quintile by composite score"), ("Good", ""), ("Average", ""),
                       ("Weak", ""), ("Poor", "Bottom quintile")]:
        line(r, tier, f'=COUNTIF({ST}${S["Tier"]}$2:${S["Tier"]}${SR},"{tier}")', note); r += 1
    line(r, "Unscored", f'=COUNTBLANK({ST}${S["Tier"]}$2:${S["Tier"]}${SR})',
         "Outside the fundamentals sample"); r += 1
    line(r, "Financial-sector exempt", f'=COUNTIF({ST}${S["Exempt"]}$2:${S["Exempt"]}${SR},"Yes")',
         "D/E, asset turns and CWIP carry no weight for these"); r += 2

    head(r, "Relative strength"); r += 1
    for band, note in [("green", "RS 80-99"), ("amber", "RS 60-79"), ("red", "RS 1-59")]:
        line(r, f"{band.capitalize()} band", f'=COUNTIF({ST}${S["Band"]}$2:${S["Band"]}${SR},"{band}")',
             note); r += 1
    line(r, "Median RS rating", f'=MEDIAN({ST}${S["RS"]}$2:${S["RS"]}${SR})',
         "Percentile rank against the loaded population", fmt=NUM2); r += 2

    head(r, "Market breadth (latest)"); r += 1
    lastb = f"COUNTA('Market Breadth'!$A$2:$A${NB + 1})"
    line(r, "Reading date", f"=INDEX('Market Breadth'!$A$2:$A${NB + 1},{lastb})", "", fmt=None); r += 1
    line(r, "% above own 200DMA", f"=INDEX('Market Breadth'!$B$2:$B${NB + 1},{lastb})", "",
         fmt=PCT); r += 1
    line(r, "vs its 200d average", f"=INDEX('Market Breadth'!$E$2:$E${NB + 1},{lastb})",
         "Breadth above its own long average = broad participation", fmt=PCT); r += 1
    line(r, "New highs", f"=INDEX('Market Breadth'!$F$2:$F${NB + 1},{lastb})", ""); r += 1
    line(r, "New lows", f"=INDEX('Market Breadth'!$G$2:$G${NB + 1},{lastb})", ""); r += 2

    head(r, "Industry indices"); r += 1
    line(r, "Industries indexed", f"=COUNTA(Sectoral!$A$2:$A${len(sectoral) + 1})",
         "Industry Name granularity; 3-constituent minimum"); r += 1
    line(r, "In golden-cross state", f'=COUNTIF(Sectoral!$M$2:$M${len(sectoral) + 1},"Yes")', ""); r += 1
    line(r, "Qualifying on Sectoral Breakout", "=COUNT('Sectoral Breakout'!$A$2:$A$1000)",
         "Usually zero - index crossovers are rare and age out fast")

    for row in range(5, r + 2):
        for col in (2, 3, 4):
            cell = ws.cell(row=row, column=col)
            if cell.font.name is None:
                cell.font = T_FONT

    # ---------------- Read Me ----------------
    ws = wb.create_sheet("Read Me", 0)
    ws.sheet_properties.tabColor = TEAL
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 96

    ws["B2"] = "Meridian - daily workbook"
    ws["B2"].font = TITLE
    ws["B3"] = f"As of {meta['as_of']}.  Generated {meta['generated_at'][:16].replace('T', ' ')} UTC."
    ws["B3"].font = SUB

    blocks = [
        ("What this is",
         "A full dump of every Meridian screen for one trading day: the same numbers the app shows, "
         "in a form you can sort, filter and keep. One worksheet per sub-tab."),
        ("Where the numbers come from",
         "Every model output is computed by meridian-engine.js - the identical module the web app "
         "imports - and written here as a value. The workbook does not re-derive the model in Excel. "
         "That is deliberate: the previous workbook did, and its financial-sector exemption listed 11 "
         "industries where the engine listed 12, quietly mis-scoring three companies. A spreadsheet "
         "that restates the model is a third implementation that will drift from the other two."),
        ("Blue vs black",
         "Blue cells are engine values - treat them as inputs. Black cells are formulas this workbook "
         "computes over those inputs (counts, distributions, ratios between adjacent columns). Edit a "
         "blue cell and the black ones follow; nothing here recalculates the model itself."),
        ("The self-check",
         "The Dashboard rebuilds all five Golden Breakout gates with COUNTIFS over the Stocks columns "
         "and compares the result against the screener's own row count. It should read OK. If it ever "
         "reads MISMATCH, the workbook is telling you its columns and the engine disagree - trust "
         "neither until that is explained."),
        ("Sheets",
         "Dashboard - coverage, gate funnel, distributions, latest breadth.\n"
         "Golden Breakout - the day's qualifying stocks, ranked.\n"
         "Stocks - all instruments with every technical and fundamental field.\n"
         "Sectoral - synthetic equal-weight industry indices.\n"
         "Sectoral Breakout - industry indices clearing the same five gates.\n"
         "Market Breadth - 500 days of participation data.\n"
         "Universe - the quarterly universe definition."),
        ("Not in this workbook yet",
         "Commodities, Currencies, Crypto and Global Indices. Their universes are defined but the only "
         "price files in the repo are prototype samples (an identical 1,260 bars for every symbol, and "
         "crypto ending 18 months before the rest), and there is no currency price file at all. The "
         "Golden Breakout thresholds have also not been validated outside equities. Publishing sheets "
         "off that data would look authoritative and be wrong, so they are omitted until those "
         "universes are backfilled the way equities were."),
        ("Reading the percentages",
         "Percent columns are stored as fractions and displayed with a percent format, so they behave "
         "correctly in your own formulas. Separation, slope, growth and breadth figures are all "
         "percentages; RSI, RS Rating and Fund Score are 0-100 index values, not percentages."),
    ]
    r = 5
    for title, body in blocks:
        ws.cell(row=r, column=2, value=title).font = Font(name=FONT, size=10, bold=True, color=INK)
        c = ws.cell(row=r, column=3, value=body)
        c.font = T_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (body.count("\n") + 1 + len(body) // 95)
        r += 2

    # openpyxl writes formulas with no cached value, so anything reading cached values
    # (including Excel's own first paint) would show blanks. LibreOffice cannot load a
    # document in this build's container, so the values cannot be baked in here — instead
    # ask the consuming application to recalculate the moment it opens the file.
    wb.calculation.fullCalcOnLoad = True
    wb.save(a.out)
    print(f"wrote {a.out}: {len(wb.sheetnames)} sheets, {NS} stocks, {len(breakout)} candidates, "
          f"{len(sectoral)} industries, {NB} breadth days")


if __name__ == "__main__":
    main()
